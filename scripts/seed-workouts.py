# -*- coding: utf-8 -*-
import json
import os
import re
from collections import OrderedDict

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
XLSX = os.path.join(ROOT, 'src', 'app', 'core', 'storage', 'Workouts.xlsx')
OUT_DIR = os.path.join(ROOT, 'src', 'assets', 'data')
OUT_JSON = os.path.join(OUT_DIR, 'seed.json')

ACTIVE_PROGRAM_ID = 2

SXR_RE = re.compile(r'^(\d+)\s*[xX]\s*(.+)$')
REST_RE = re.compile(r'^(\d+)\s*a\s*(\d+)\s*s$', re.IGNORECASE)


def parse_sxr(value):
    raw = (value or '').strip()
    match = SXR_RE.match(raw)
    if not match:
        return raw, None, raw
    return raw, int(match.group(1)), match.group(2).strip()


def parse_rest(value):
    raw = (value or '').strip()
    match = REST_RE.match(raw)
    if not match:
        return raw, None, None
    return raw, int(match.group(1)), int(match.group(2))


def iso_date(value):
    if hasattr(value, 'strftime'):
        return value.strftime('%Y-%m-%d')
    return str(value)[:10]


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    programs = OrderedDict()
    exercises = OrderedDict()
    workouts = OrderedDict()
    workout_exercises = []
    workout_id_seq = 0
    exercise_id_seq = 0
    we_id_seq = 0
    sort_by_workout = {}
    workouts_in_program = {}

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        program_id, date, letter, muscle, exercise_name, sxr, technique, rest, program_type = row
        if program_id is None:
            continue

        program_id = int(program_id)
        letter = (letter or '').strip()
        muscle = (muscle or '').strip()
        exercise_name = (exercise_name or '').strip()
        started_on = iso_date(date)

        if program_id not in programs:
            programs[program_id] = {
                'id': program_id,
                'startedOn': started_on,
                'type': (program_type or '').strip(),
                'isActive': 1 if program_id == ACTIVE_PROGRAM_ID else 0
            }

        workout_key = (program_id, letter)
        if workout_key not in workouts:
            workout_id_seq += 1
            workouts_in_program[program_id] = workouts_in_program.get(program_id, 0) + 1
            workouts[workout_key] = {
                'id': workout_id_seq,
                'programId': program_id,
                'letter': letter,
                'targetMuscle': muscle,
                'sortOrder': workouts_in_program[program_id]
            }

        if exercise_name not in exercises:
            exercise_id_seq += 1
            exercises[exercise_name] = {
                'id': exercise_id_seq,
                'name': exercise_name
            }

        sxr_raw, sets, reps = parse_sxr(sxr)
        rest_raw, rest_min, rest_max = parse_rest(rest)
        workout_id = workouts[workout_key]['id']
        sort_by_workout[workout_id] = sort_by_workout.get(workout_id, 0) + 1
        we_id_seq += 1
        workout_exercises.append({
            'id': we_id_seq,
            'workoutId': workout_id,
            'exerciseId': exercises[exercise_name]['id'],
            'sortOrder': sort_by_workout[workout_id],
            'sxr': sxr_raw,
            'sets': sets,
            'reps': reps,
            'technique': (technique or '').strip() or None,
            'restPeriod': rest_raw,
            'restMinSec': rest_min,
            'restMaxSec': rest_max
        })

    seed = {
        'programs': list(programs.values()),
        'workouts': list(workouts.values()),
        'exercises': list(exercises.values()),
        'workoutExercises': workout_exercises
    }

    os.makedirs(OUT_DIR, exist_ok=True)
    with open(OUT_JSON, 'w', encoding='utf-8') as handle:
        json.dump(seed, handle, ensure_ascii=False, separators=(',', ':'))

    print('programs', len(seed['programs']))
    print('workouts', len(seed['workouts']))
    print('exercises', len(seed['exercises']))
    print('workoutExercises', len(seed['workoutExercises']))
    print('wrote', OUT_JSON)


if __name__ == '__main__':
    main()
